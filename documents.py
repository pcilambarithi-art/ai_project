from io import BytesIO
from pathlib import Path
from PIL import Image
from pypdf import PdfReader
from docx import Document
import pytesseract

ALLOWED = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".webp": "image/webp", ".bmp": "image/bmp",
}
MAX_TEXT = 120_000

def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED:
        raise ValueError(f"Unsupported file type: {ext or 'unknown'}")
    if not data:
        return ""
    if ext == ".txt":
        result = data.decode("utf-8-sig", errors="replace")
    elif ext == ".pdf":
        reader = PdfReader(BytesIO(data))
        result = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == ".docx":
        doc = Document(BytesIO(data))
        result = "\n".join([p.text for p in doc.paragraphs] + [
            " | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows
        ])
    else:
        try:
            result = pytesseract.image_to_string(Image.open(BytesIO(data)))
        except Exception as exc:
            raise ValueError("Image OCR failed. Install Tesseract OCR and ensure it is on PATH.") from exc
    return result[:MAX_TEXT]
